# -*- coding: utf-8 -*-
"""
Gera o kickoff EquiDados em .pptx usando python-pptx.
Rodar com: pip install python-pptx && python gerar_kickoff_pptx.py
"""

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

DESTINO = os.path.expanduser("~/Documents/Kickoff_EquiDados.pptx")

DARK = RGBColor(0x23, 0x23, 0x23)
ORANGE = RGBColor(0xF4, 0x61, 0x1E)
CREAM = RGBColor(0xFC, 0xF3, 0xE6)
WHITE = RGBColor(0xFF, 0xF8, 0xEF)
TEXT_DARK = RGBColor(0x23, 0x23, 0x23)

SLIDE_W = Emu(12192000)  # 13.333in
SLIDE_H = Emu(6858000)   # 7.5in

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


_slide_count = [0]


def add_slide():
    _slide_count[0] += 1
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_circle(slide, x, y, size, fill=None, outline=None, outline_w=2.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x), Emu(y), Emu(size), Emu(size))
    shape.shadow.inherit = False
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if outline is not None:
        shape.line.color.rgb = outline
        shape.line.width = Pt(outline_w)
    else:
        shape.line.fill.background()
    return shape


def add_bottom_bar(slide, color=ORANGE, height=140000):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), Emu(SLIDE_H - height), SLIDE_W, Emu(height)
    )
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_textbox(slide, x, y, w, h, runs, anchor=MSO_ANCHOR.TOP):
    """runs: list of paragraphs, each a list of (text, size, bold, color, bullet, italic)"""
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for para_runs in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(10)
        for (text, size, bold, color, bullet, italic) in para_runs:
            if bullet:
                p.text = ""  # bullets handled via manual "•" prefix below
            r = p.add_run()
            r.text = ("•  " + text) if bullet else text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
            r.font.name = "Arial"
    return box


def run(text, size=18, bold=False, color=TEXT_DARK, bullet=False, italic=False):
    return [(text, size, bold, color, bullet, italic)]


def title_slide(titulo, subtitulo, corpo):
    s = add_slide()
    set_bg(s, DARK)
    add_circle(s, SLIDE_W - 1900000, -700000, 2600000, fill=ORANGE)
    add_circle(s, SLIDE_W - 700000, SLIDE_H - 1600000, 2000000, fill=ORANGE)
    add_circle(s, -500000, SLIDE_H - 900000, 1600000, fill=CREAM)
    add_textbox(s, 900000, 2000000, 9600000, 1300000, [run(titulo, 60, True, ORANGE, italic=True)])
    add_textbox(s, 900000, 3350000, 9600000, 700000, [run(subtitulo, 26, True, WHITE)])
    add_textbox(s, 900000, 4200000, 9800000, 900000, [run(corpo, 18, False, WHITE)])
    return s


def cream_slide(titulo, blocks, dark_accent=True):
    """blocks: list of ('h'|'p'|'b', text, size)"""
    s = add_slide()
    set_bg(s, CREAM)
    add_circle(s, SLIDE_W - 1400000, -600000, 1900000, fill=ORANGE)
    add_circle(s, SLIDE_W - 300000, 900000, 900000, fill=(DARK if dark_accent else ORANGE))
    add_bottom_bar(s)
    add_textbox(s, 700000, 450000, 9500000, 900000, [run(titulo, 34, True, ORANGE, italic=True)])

    paras = []
    for kind, text, *rest in blocks:
        size = rest[0] if rest else (20 if kind == "h" else 17)
        if kind == "h":
            paras.append(run(text, size, True, ORANGE))
        elif kind == "b":
            paras.append(run(text, size, False, TEXT_DARK, bullet=True))
        elif kind == "note":
            paras.append(run(text, size, False, RGBColor(0x6B, 0x64, 0x59), italic=True))
        else:
            paras.append(run(text, size, False, TEXT_DARK))
    add_textbox(s, 700000, 1500000, 10800000, 4900000, paras)
    return s


def dark_section_slide(titulo, blocks):
    s = add_slide()
    set_bg(s, DARK)
    add_circle(s, SLIDE_W - 1700000, -600000, 2300000, fill=ORANGE)
    add_circle(s, -600000, SLIDE_H - 1100000, 1800000, fill=ORANGE)
    add_textbox(s, 900000, 900000, 9800000, 900000, [run(titulo, 40, True, ORANGE, italic=True)])
    paras = []
    for kind, text, *rest in blocks:
        size = rest[0] if rest else (20 if kind == "h" else 17)
        if kind == "h":
            paras.append(run(text, size, True, WHITE))
        else:
            paras.append(run(text, size, False, WHITE))
    add_textbox(s, 900000, 2000000, 9800000, 4200000, paras)
    return s


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

# 1
title_slide(
    "EquiDados",
    "Monitoramento de Equidade em Saúde — Recife",
    "Fabiana Lima  ·  Felipe Saraiva Belém  ·  José Bruno  ·  Rodrigo Aguiar  ·  Anna Clara",
)

# 2
cream_slide("Roteiro", [
    ("b", "Introdução e contexto do desafio"),
    ("b", "Problema e justificativa"),
    ("b", "Objetivos do projeto"),
    ("b", "Análises: Matriz CSD, Personas, Mapa de Empatia, SWOT"),
    ("b", "Processo de ideação"),
    ("b", "Proposta de solução: pipeline de dados, protótipos e Machine Learning"),
    ("b", "Cronograma e encerramento"),
])

# 3
cream_slide("Introdução — Contexto", [
    ("p", "A Secretaria de Atenção Básica do Recife conduz três políticas municipais de "
          "equidade em saúde: Saúde Integral da População LGBTQIAPN+, Atenção Integral à "
          "Saúde da Pessoa com Deficiência e Saúde Integral da População Negra.", 19),
    ("p", "Os dados necessários para monitorar essas políticas já existem no sistema "
          "PEC/e-SUS APS — identidade de gênero, orientação sexual, nome social, raça/cor, "
          "tipo de deficiência.", 19),
])

# 4
cream_slide("Introdução — Problemática", [
    ("b", "Dados dispersos em diferentes campos e registros do PEC/e-SUS"),
    ("b", "Processo hoje manual, pontual e não padronizado de extração e análise"),
    ("b", "Baixa visibilidade sobre o perfil territorial e sociodemográfico das três populações"),
    ("b", "Fragilidade específica no preenchimento do quesito raça/cor — obrigatório, mas sem "
          "monitoramento de completude ou consistência"),
])

# 5
dark_section_slide("Problema", [
    ("p", "A gestão não consegue monitorar, de forma automatizada e territorializada, a "
          "situação das três populações de equidade em saúde.", 22),
    ("p", "Os dados existem — mas estão dispersos e sem padronização, exigindo hoje um "
          "processo manual a cada solicitação de relatório.", 22),
])

# 6
cream_slide("Justificativa", [
    ("p", "Políticas públicas de equidade racial, de gênero e de acessibilidade dependem de "
          "evidência para priorizar ações e recursos.", 20),
    ("p", "Sem dados estruturados e territorializados, desigualdades entre Distritos "
          "Sanitários passam despercebidas e a gestão permanece reativa em vez de proativa.", 20),
])

# 7
cream_slide("Objetivos", [
    ("h", "Objetivo geral"),
    ("p", "Construir um pipeline automatizado que extraia, integre e disponibilize os dados "
          "das três populações de equidade em indicadores territoriais.", 17),
    ("h", "Objetivos específicos"),
    ("b", "Automatizar a extração de fontes públicas e, futuramente, do PEC/e-SUS", 16),
    ("b", "Estruturar os dados em camadas (bronze e silver)", 16),
    ("b", "Disponibilizar indicadores por Distrito Sanitário, Unidade e Equipe de Saúde", 16),
    ("b", "Monitorar a qualidade do preenchimento dos campos estratégicos", 16),
])

# 8
cream_slide("Análises — Matriz CSD", [
    ("h", "Certezas"),
    ("b", "Os dados já existem no PEC/e-SUS; o problema é a falta de automação", 16),
    ("b", "Três públicos-alvo definidos: LGBTQIAPN+, PCD e população negra", 16),
    ("h", "Suposições"),
    ("b", "O objetivo final é um painel de consulta contínua, não um relatório único", 16),
    ("h", "Dúvidas em aberto"),
    ("b", "Priorização entre os três públicos-alvo", 16),
    ("b", "Forma de acesso aos dados do PEC/e-SUS", 16),
    ("b", "Restrições legais sobre armazenamento de dados sensíveis", 16),
])

# 9
cream_slide("Análises — Personas e Mapa de Empatia", [
    ("h", "Camila — Gestora"),
    ("p", "Precisa de um painel de fácil leitura, atualizado, por território e público-alvo.", 16),
    ("h", "Rafael — Técnico de Dados"),
    ("p", "Precisa de um pipeline automatizado com alertas de qualidade de cadastro.", 16),
    ("h", "Josiane — Equipe de Saúde da Família"),
    ("p", "Precisa de orientação clara sobre como coletar dados sensíveis com respeito.", 16),
])

# 10
cream_slide("Análises — SWOT", [
    ("h", "Forças"),
    ("b", "Pipeline de ETL funcional ponta a ponta; agente de consulta local sem custo de API", 15),
    ("h", "Fraquezas"),
    ("b", "Indicadores centrais dependem do acesso ao PEC/e-SUS, ainda não concedido", 15),
    ("b", "Nenhuma fonte pública hoje cobre orientação sexual ou identidade de gênero", 15),
    ("h", "Oportunidades"),
    ("b", "Potencial de servir como piloto replicável para outras secretarias", 15),
    ("h", "Ameaças"),
    ("b", "Liberação do PEC/e-SUS pode não ocorrer dentro do prazo do desafio", 15),
    ("note", "Benchmarking (Estudo de Similares) em consolidação — previsto para as próximas entregas.", 13),
])

# 11
cream_slide("Processo de Ideação", [
    ("p", "As técnicas de ideação (Brainstorming, Brainwriting, SCAMPER, Crazy 8's) estão em "
          "consolidação ao longo desta semana.", 19),
    ("p", "A proposta de solução já validada até aqui — pipeline de dados em camadas, agente "
          "de consulta em linguagem natural e modelo de classificação — nasceu da análise "
          "direta do problema levantado junto à Secretaria, e será formalizada com as "
          "técnicas de ideação nas próximas entregas.", 19),
])

# 12
cream_slide("Proposta de Solução — Pipeline de Dados", [
    ("b", "Bronze: extração automatizada (CNES, IBGE, Dados Recife), armazenada em MongoDB", 17),
    ("b", "Silver: limpeza e padronização em banco relacional (Neon/PostgreSQL)", 17),
    ("b", "Consumo: painéis no Metabase e agente de consulta em linguagem natural (LLM local)", 17),
    ("h", "Diferencial já validado"),
    ("p", "Achado real de qualidade de dados: identificação de unidades de saúde sem Distrito "
          "Sanitário preenchido no cadastro oficial — evidência de que o método funciona antes "
          "mesmo da liberação do PEC.", 15),
])

# 13
cream_slide("Proposta de Solução — Protótipos (Wireframes)", [
    ("b", "Painel de Equidade em Saúde — visão da gestora, indicadores por território", 18),
    ("b", "Monitor de Qualidade do Cadastro — visão técnica, ranking de unidades críticas", 18),
    ("b", "Painel de Insights — incorporação de Machine Learning e LLM (camada Gold)", 18),
    ("b", "Assistente Equidados IA — consulta em linguagem natural sobre os dados", 18),
])

# 14
cream_slide("Proposta de Solução — Machine Learning", [
    ("h", "EquiDados: Sistema de Inteligência em Equidade em Saúde"),
    ("p", "Populações vulnerabilizadas não têm acesso adequado porque o sistema não identifica "
          "automaticamente as adaptações e serviços que cada uma demanda (ex.: interpretação "
          "em Libras, endocrinologista de gênero, psicólogo sensível a questões raciais).", 16),
    ("p", "Solução: modelo de Classificação Multilabel que, a partir do histórico do paciente "
          "no e-SUS APS/PEC, prediz simultaneamente múltiplas necessidades de adaptação e "
          "cuidado — reconhecendo inclusive interseccionalidade (ex.: mulher negra trans com "
          "deficiência).", 16),
])

# 15
cream_slide("Machine Learning — Features, Target e Impacto", [
    ("h", "Features (e-SUS APS/PEC)"),
    ("b", "Raça/cor, gênero registrado, orientação sexual, tipo e grau de deficiência", 15),
    ("b", "Histórico de consultas, faltas e adaptações solicitadas; bairro de residência", 15),
    ("h", "Target"),
    ("b", "Variáveis binárias de necessidade (ex.: necessidade_interprete_libras, "
          "necessidade_psicologo_LGBTQ, necessidade_atendimento_antidiscriminatorio)", 15),
    ("h", "Impacto prático"),
    ("b", "Equipes preparadas antes da consulta; gestão com visibilidade real de demanda; "
          "pacientes com acesso garantido e atendimento digno", 15),
])

# 16
dark_section_slide("Cronograma e Encerramento", [
    ("h", "Concluído"),
    ("p", "Mapeamento de fontes públicas, pipeline bronze/silver funcional, agente de consulta "
          "local, documentação inicial (CSD, personas, wireframes).", 16),
    ("h", "Próximos passos"),
    ("p", "Validação da Matriz CSD com a Secretaria, construção do painel de indicadores, "
          "solicitação formal de acesso ao PEC/e-SUS.", 16),
])

os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
prs.save(DESTINO)
print("Salvo em:", DESTINO)
print("Total de slides:", _slide_count[0])
